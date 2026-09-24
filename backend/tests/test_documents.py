# ruff: noqa: F811  (pytest fixture imported from test_amocrm)
"""Documents from templates for amoCRM leads: Word fill, numbering, widget, site."""

import io
import json

import httpx
import pytest
import respx
from docx import Document

from documents import render, services
from documents.models import DocTemplate, GeneratedDoc
from sms.models import SmsVariable
from tests.test_amocrm import HOST, conn, widget_token  # noqa: F401

API = {"HTTP_HOST": "api.uzbridge.test"}
SITE = {"HTTP_HOST": "acme.uzbridge.test"}


def word(*paragraphs, split: str = "") -> bytes:
    """A .docx; `split` puts that keyword across two runs, as Word often does."""
    doc = Document()
    for text in paragraphs:
        p = doc.add_paragraph()
        if split and split in text:
            before, after = text.split(split, 1)
            cut = len(split) // 2
            p.add_run(before)
            bold = p.add_run(split[:cut])
            bold.bold = True
            p.add_run(split[cut:] + after)
        else:
            p.add_run(text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Mijoz"
    table.cell(0, 1).text = "{name}"
    doc.sections[0].header.paragraphs[0].text = "№ {number}"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def text_of(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for row in doc.tables[0].rows for c in row.cells]
    parts += [p.text for p in doc.sections[0].header.paragraphs]
    return "\n".join(parts)


def test_word_fill_keeps_formatting_and_joins_split_keywords():
    data = word("Shartnoma {number} sana {date}", "Summa: {amount} ({amount_words}) {unknown}", split="{number}")
    out = render.fill_docx(
        data,
        {
            "number": "DOG-0007",
            "date": "25.09.2026",
            "amount": "1 500 000",
            "name": "Ali",
            "amount_words": services.amount_in_words(1_500_000),
        },
    )
    text = text_of(out)
    assert "Shartnoma DOG-0007 sana 25.09.2026" in text
    assert "1 500 000 (bir million besh yuz ming)" in text
    assert "{unknown}" in text  # an unknown keyword stays visible
    assert "Ali" in text and "№ DOG-0007" in text  # tables and header too
    # the run that held the start of the keyword keeps its bold
    runs = Document(io.BytesIO(out)).paragraphs[0].runs
    assert any(r.bold and "DOG-0007" in r.text for r in runs)
    assert render.docx_keywords(data)[:2] == ["number", "date"]


def test_amount_in_words():
    assert services.amount_in_words(0) == "nol"
    assert services.amount_in_words(115) == "bir yuz oʻn besh"
    assert services.amount_in_words(2_000_450) == "ikki million toʻrt yuz ellik"


def mock_lead(price=1_500_000):
    respx.get(f"https://{HOST}/api/v4/leads/9001").mock(
        return_value=httpx.Response(
            200,
            json={
                "id": 9001,
                "name": "Olma savdo",
                "price": price,
                "_embedded": {"contacts": [{"id": 5, "is_main": True}]},
            },
        )
    )
    respx.get(f"https://{HOST}/api/v4/contacts/5").mock(
        return_value=httpx.Response(
            200,
            json={
                "name": "Ali Valiyev",
                "custom_fields_values": [{"field_code": "PHONE", "values": [{"value": "+998901234567"}]}],
            },
        )
    )
    return respx.post(f"https://{HOST}/api/v4/leads/9001/notes").mock(return_value=httpx.Response(200, json={}))


@pytest.fixture
def template(company):
    return DocTemplate.objects.create(
        company=company,
        name="Shartnoma",
        kind="docx",
        docx=word("Shartnoma {number}", "{lead_name}: {amount} soʻm"),
        default_format="docx",
        prefix="DOG-",
        padding=4,
        next_number=7,
    )


@pytest.mark.django_db
@respx.mock
def test_numbering_runs_on_and_regenerating_keeps_the_number(conn, company, template):
    note = mock_lead()
    first = services.generate(conn, template, 9001, "docx", user_label="Aziz")
    assert (first.number, first.seq) == ("DOG-0007", 7)
    assert "Shartnoma DOG-0007" in text_of(bytes(first.content))
    assert "Olma savdo: 1 500 000 soʻm" in text_of(bytes(first.content))
    assert "DOG-0007" in json.loads(note.calls[0].request.content)[0]["params"]["text"]

    again = services.generate(conn, template, 9001, "docx")
    assert again.pk == first.pk and again.number == "DOG-0007"  # same lead: same number
    template.refresh_from_db()
    assert template.next_number == 8

    respx.get(f"https://{HOST}/api/v4/leads/9002").mock(return_value=httpx.Response(200, json={"id": 9002, "price": 0}))
    respx.post(f"https://{HOST}/api/v4/leads/9002/notes").mock(return_value=httpx.Response(200, json={}))
    other = services.generate(conn, template, 9002, "docx")
    assert other.number == "DOG-0008"


@pytest.mark.django_db
@respx.mock
def test_widget_generates_only_when_switched_on(client, conn, company, template, monkeypatch):
    mock_lead()
    kw = {"HTTP_X_AUTH_TOKEN": widget_token(), **API}
    body = {"lead_id": 9001, "template_id": template.pk, "format": "pdf"}
    r = client.post("/api/widget/docs/generate", data=json.dumps(body), content_type="application/json", **kw)
    assert r.status_code == 409  # off by default

    conn.docs_enabled = True
    conn.save()
    monkeypatch.setattr(render, "docx_to_pdf", lambda data: b"%PDF-1.7 fake")
    r = client.post("/api/widget/docs/generate", data=json.dumps(body), content_type="application/json", **kw)
    assert r.status_code == 200, r.content
    made = r.json()["made"]
    assert made["number"] == "DOG-0007" and made["format"] == "pdf"
    listed = client.get("/api/widget/docs?lead_id=9001", **kw).json()
    assert [d["number"] for d in listed["documents"]] == ["DOG-0007"]
    assert listed["templates"][0]["next"] == "DOG-0008"

    # the note's link downloads the file without signing in
    token = made["url"].rstrip("/").split("/")[-1]
    file = client.get(f"/d/{token}/", **API)
    assert file.status_code == 200 and file["Content-Type"] == "application/pdf" and file.content.startswith(b"%PDF")
    assert client.get("/d/not-a-token/", **API).status_code == 404

    template.use_in_amocrm = False
    template.save()
    r = client.post("/api/widget/docs/generate", data=json.dumps(body), content_type="application/json", **kw)
    assert r.status_code == 404


@pytest.mark.django_db
@respx.mock
def test_site_templates_upload_keywords_and_preview(client, conn, company, owner, monkeypatch):
    client.force_login(owner)
    csrf = client.get("/api/auth/csrf", **SITE).cookies["csrftoken"].value
    kw = {**SITE, "HTTP_X_CSRFTOKEN": csrf}
    SmsVariable.objects.create(company=company, integration="amocrm", key="muddat", source="lead.cf.501")

    # built on the site
    r = client.post(
        "/api/documents/templates",
        data=json.dumps(
            {"name": "Dalolatnoma", "kind": "html", "html": "<h1>{number}</h1><p>{muddat} {typo}</p>", "prefix": "AKT-"}
        ),
        content_type="application/json",
        **kw,
    )
    assert r.status_code == 200, r.content
    html_t = r.json()
    assert html_t["keywords"] == ["number", "muddat", "typo"] and html_t["unknown_keywords"] == ["typo"]
    assert html_t["next_label"] == "AKT-0001"

    # uploaded Word
    word_t = client.post(
        "/api/documents/templates",
        data=json.dumps({"name": "Shartnoma", "kind": "docx", "default_format": "docx"}),
        content_type="application/json",
        **kw,
    ).json()
    bad = client.post(f"/api/documents/templates/{word_t['id']}/upload", {"file": io.BytesIO(b"not word")}, **kw)
    assert bad.status_code == 422
    upload = io.BytesIO(word("Shartnoma {number} {name}"))
    upload.name = "shartnoma.docx"
    r = client.post(f"/api/documents/templates/{word_t['id']}/upload", {"file": upload}, **kw)
    assert r.status_code == 200, r.content
    assert r.json()["has_file"] and "number" in r.json()["keywords"]

    # the counter can continue an old numbering, but never go back under used numbers
    r = client.put(
        f"/api/documents/templates/{word_t['id']}",
        data=json.dumps(
            {"name": "Shartnoma", "kind": "docx", "default_format": "docx", "prefix": "DOG-", "next_number": 120}
        ),
        content_type="application/json",
        **kw,
    )
    assert r.json()["next_label"] == "DOG-0120"

    # preview fills a real lead but uses no number
    mock_lead()
    r = client.post(
        f"/api/documents/templates/{word_t['id']}/preview",
        data=json.dumps({"lead_id": 9001, "format": "docx"}),
        content_type="application/json",
        **kw,
    )
    assert r.status_code == 200, r.content
    assert "DOG-0120" in text_of(r.content) and "Ali Valiyev" in text_of(r.content)
    assert DocTemplate.objects.get(pk=word_t["id"]).next_number == 120
    assert not GeneratedDoc.objects.exists()

    # HTML templates go through LibreOffice; here it's stood in for
    monkeypatch.setattr(render, "html_to", lambda fmt, page: page.encode())
    r = client.post(
        f"/api/documents/templates/{html_t['id']}/preview",
        data=json.dumps({"lead_id": 9001, "format": "pdf"}),
        content_type="application/json",
        **kw,
    )
    assert r.status_code == 200 and b"<h1>AKT-0001</h1>" in r.content and b"{typo}" in r.content

    assert client.get("/api/documents/keywords", **SITE).json()["own"][0]["key"] == "muddat"


DRIVE = "https://drive-b.amocrm.ru"


def mock_drive():
    respx.get(f"https://{HOST}/api/v4/account").mock(
        return_value=httpx.Response(200, json={"id": 31337, "drive_url": DRIVE})
    )
    session = respx.post(f"{DRIVE}/v1.0/sessions").mock(
        return_value=httpx.Response(
            200, json={"session_id": 1, "upload_url": f"{DRIVE}/v1.0/sessions/upload/tok1", "max_part_size": 524288}
        )
    )
    part = respx.post(f"{DRIVE}/v1.0/sessions/upload/tok1").mock(
        return_value=httpx.Response(200, json={"uuid": "file-uuid-1", "name": "x.docx"})
    )
    attach = respx.put(f"https://{HOST}/api/v4/leads/9001/files").mock(return_value=httpx.Response(200, json={}))
    return session, part, attach


@pytest.mark.django_db
@respx.mock
def test_keyword_sources_per_template_and_files_tab(conn, company, freezer=None):
    from datetime import date
    from unittest.mock import patch

    t = DocTemplate.objects.create(
        company=company,
        name="Shartnoma",
        kind="docx",
        docx=word("{raqam} | {muddat} | {sana} | {shahar} | {stir}"),
        default_format="docx",
        prefix="SH-",
        bindings=services.clean_bindings(
            {
                "raqam": {"type": "number"},
                "muddat": {"type": "amo", "source": "lead.cf.501"},
                "sana": {"type": "date", "offset_days": 10, "style": "long"},
                "shahar": {"type": "text", "value": "Toshkent"},
                "stir": {"type": "builtin", "key": "company_tin"},
            }
        ),
    )
    company.tin = "123456789"
    company.save()
    respx.get(f"https://{HOST}/api/v4/leads/9001").mock(
        return_value=httpx.Response(
            200,
            json={
                "id": 9001,
                "price": 1000,
                "custom_fields_values": [{"field_id": 501, "values": [{"value": "30 kun"}]}],
                "_embedded": {"contacts": []},
            },
        )
    )
    respx.post(f"https://{HOST}/api/v4/leads/9001/notes").mock(return_value=httpx.Response(200, json={}))
    session, part, attach = mock_drive()
    with patch("documents.services.timezone.localdate", return_value=date(2026, 9, 25)):
        doc = services.generate(conn, t, 9001, "docx")
    assert "SH-0001 | 30 kun | 5-oktyabr 2026-yil | Toshkent | 123456789" in text_of(bytes(doc.content))
    # in the lead's "Файлы" tab
    assert json.loads(session.calls[0].request.content)["file_name"] == doc.filename
    assert part.calls[0].request.content == bytes(doc.content)
    assert json.loads(attach.calls[0].request.content) == [{"file_uuid": "file-uuid-1"}]
    doc.refresh_from_db()
    assert (doc.amo_file_uuid, doc.amo_file_format) == ("file-uuid-1", "docx")
    # re-made in the same format: a new version of the same file, not a second file
    services.generate(conn, t, 9001, "docx")
    assert json.loads(session.calls[1].request.content)["file_uuid"] == "file-uuid-1"
    assert len(attach.calls) == 1

    with pytest.raises(ValueError):
        services.clean_bindings({"x": {"type": "amo", "source": "lead.cf.abc"}})


@pytest.mark.django_db
@respx.mock
def test_no_file_access_still_links_in_a_note(conn, company, template):
    note = mock_lead()
    respx.get(f"https://{HOST}/api/v4/account").mock(return_value=httpx.Response(200, json={"drive_url": DRIVE}))
    respx.post(f"{DRIVE}/v1.0/sessions").mock(return_value=httpx.Response(403, json={"title": "Forbidden"}))
    doc = services.generate(conn, template, 9001, "docx")
    assert doc.amo_file_uuid == ""
    assert "/d/" in json.loads(note.calls[-1].request.content)[0]["params"]["text"]
    conn.refresh_from_db()
    assert conn.status == "active"  # a missing files scope doesn't break the connection


@pytest.mark.django_db
def test_word_template_can_be_edited_on_the_site_and_back(client, company, owner, template):
    client.force_login(owner)
    csrf = client.get("/api/auth/csrf", **SITE).cookies["csrftoken"].value
    kw = {**SITE, "HTTP_X_CSRFTOKEN": csrf}
    r = client.post(f"/api/documents/templates/{template.pk}/edit-on-site", **kw)
    assert r.status_code == 200, r.content
    body = r.json()
    assert body["kind"] == "html" and "{number}" in body["html"] and body["can_edit_word"]
    back = client.post(f"/api/documents/templates/{template.pk}/use-word", **kw).json()
    assert back["kind"] == "docx"


@pytest.mark.django_db
@respx.mock
def test_documents_is_a_paid_switch(client, conn, company):
    from accounts.models import Company
    from billing.models import LedgerEntry

    respx.get(f"https://{HOST}/api/v4/users/42").mock(
        return_value=httpx.Response(200, json={"rights": {"is_admin": True}})
    )
    respx.get(f"https://{HOST}/api/v4/leads/custom_fields").mock(
        return_value=httpx.Response(200, json={"_embedded": {"custom_fields": []}})
    )
    respx.post(f"https://{HOST}/api/v4/leads/custom_fields").mock(
        return_value=httpx.Response(200, json={"_embedded": {"custom_fields": [{"id": 1}]}})
    )
    respx.route(host=HOST, path__startswith="/api/v4/webhooks").mock(return_value=httpx.Response(200, json={}))
    kw = {"HTTP_X_AUTH_TOKEN": widget_token(), **API}

    def switch(on):
        body = {"accounts": [], "sms_enabled": True, "docs_enabled": on}
        return client.put("/api/widget/settings", data=json.dumps(body), content_type="application/json", **kw)

    Company.objects.filter(pk=company.pk).update(balance_tiyin=0)
    r = switch(True)
    assert r.status_code == 402 and "Documents" in r.json()["detail"]
    conn.refresh_from_db()
    assert not conn.docs_enabled  # rolled back

    Company.objects.filter(pk=company.pk).update(balance_tiyin=1_000_000 * 100)
    assert switch(True).status_code == 200
    assert LedgerEntry.objects.get().amount_tiyin == -200_000 * 100
    conn.refresh_from_db()
    assert conn.docs_enabled
